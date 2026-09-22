"""Servicio de exportación de reportes a Excel / HTML / PDF (CU20/CU15).

Reutiliza los dict que devuelven ReporteService.generar_reporte_* sin tocar su lógica.
"""
import io
from datetime import datetime
from typing import Any, Dict, List, Tuple


def _fecha_tag() -> str:
    return datetime.now().strftime("%Y%m%d")


def _resumen_rows(datos: Dict[str, Any]) -> List[List[Any]]:
    rows: List[List[Any]] = []
    for key in ("periodo", "cliente", "cliente_nombre", "total_compras", "total_gastado",
                "total_pedidos_online", "total_ventas_presenciales",
                "total_reservas", "reservas_activas", "reservas_completadas", "reservas_vencidas",
                "resumen", "tasa_conversion_recogida_pct",
                "monto_total_convertido", "total_clientes_registrados",
                "clientes_activos_con_compras", "total_items_registrados",
                "total_unidades_disponibles", "items_con_bajo_stock",
                "beneficio_estimado_margen_40pct", "ingresos"):
        if key in datos:
            v = datos[key]
            if isinstance(v, dict):
                for sk, sv in v.items():
                    rows.append([f"{key}.{sk}", sv])
            else:
                rows.append([key, v])
    return rows


def _detalle_rows(tipo: str, datos: Dict[str, Any]) -> Tuple[List[str], List[List[Any]]]:
    # Compras del Cliente
    if "compras" in datos and isinstance(datos.get("compras"), list):
        head = ["Código", "Fecha", "Tipo", "Items", "Total (Bs.)", "Método Pago", "Estado"]
        body = [[
            c.get("codigo", f"#{c.get('id', '')}"),
            c.get("fecha", ""),
            c.get("tipo", "Online"),
            c.get("total_items", len(c.get("detalles", []))),
            c.get("total", 0.0),
            c.get("metodo_pago", "-"),
            c.get("estado", "")
        ] for c in datos["compras"]]
        return head, body

    # Reservas del Cliente
    if "reservas" in datos and isinstance(datos.get("reservas"), list):
        head = ["Código", "Sucursal", "Fecha Reserva", "Fecha Límite", "Prendas", "Total Estimado (Bs.)", "Estado"]
        body = [[
            r.get("codigo_reserva", f"#{r.get('id', '')}"),
            r.get("sucursal", ""),
            r.get("fecha_reserva", ""),
            r.get("fecha_limite", ""),
            r.get("total_items", len(r.get("detalles", []))),
            r.get("total_estimado", 0.0),
            r.get("estado", "")
        ] for r in datos["reservas"]]
        return head, body

    if "top_productos" in datos and datos.get("top_productos"):
        head = ["ID Producto", "Nombre", "Unidades Vendidas"]
        body = [[p.get("producto_id"), p.get("nombre"), p.get("unidades_vendidas")]
                for p in datos["top_productos"]]
        return head, body
    if "inventario" in datos and datos.get("inventario"):
        head = ["ID", "Producto", "SKU", "Talla", "Color", "Sucursal", "Disponible", "Reservado", "Mínimo", "Alerta"]
        body = [[i.get("inventario_id"), i.get("producto_nombre"), i.get("sku"),
                 i.get("talla"), i.get("color"), i.get("sucursal"),
                 i.get("cantidad_disponible"), i.get("cantidad_reservada"), i.get("cantidad_minima"),
                 "SI" if i.get("alerta_bajo_stock") else "NO"]
                for i in datos["inventario"]]
        return head, body
    if "desglose_estados" in datos and isinstance(datos["desglose_estados"], dict):
        head = ["Estado", "Cantidad"]
        body = [[k, v] for k, v in datos["desglose_estados"].items()]
        return head, body
    if "top_10_clientes" in datos and datos.get("top_10_clientes"):
        head = ["ID", "Nombre", "Email", "Pedidos", "Total Gastado"]
        body = [[c.get("cliente_id"), c.get("nombre"), c.get("email", c.get("correo")),
                 c.get("total_pedidos"), c.get("total_gastado")]
                for c in datos["top_10_clientes"]]
        return head, body
    if "desglose_por_metodo_pago" in datos and isinstance(datos["desglose_por_metodo_pago"], dict):
        head = ["Método de pago", "Total"]
        body = [[k, v] for k, v in datos["desglose_por_metodo_pago"].items()]
        return head, body
    if "kpis_detallados" in datos and isinstance(datos["kpis_detallados"], list):
        head = ["KPI", "Actual", "Objetivo", "Unidad"]
        body = [[k.get("nombre"), k.get("valor_actual"), k.get("valor_objetivo"),
                 k.get("unidad_medida")] for k in datos["kpis_detallados"]]
        return head, body
    return [], []


def exportar_excel(tipo: str, datos: Dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = tipo[:31]
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="E11D48")
    ws.append([f"Reporte {tipo}", datetime.now().strftime("%d/%m/%Y %H:%M")])
    ws.append([])
    for r in _resumen_rows(datos):
        ws.append(r)
    ws.append([])
    head, body = _detalle_rows(tipo, datos)
    if head:
        ws.append(head)
        for c in ws[ws.max_row]:
            c.font = hdr_font
            c.fill = hdr_fill
        for row in body[:200]:
            ws.append(row)
    for col in ws.columns:
        maxlen = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(48, max(12, maxlen + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def exportar_html(tipo: str, datos: Dict[str, Any]) -> str:
    head, body = _detalle_rows(tipo, datos)
    resumen = "".join(f"<tr><td>{a}</td><td>{b}</td></tr>" for a, b in _resumen_rows(datos))
    det = ""
    if head:
        det = "<h2>Detalle</h2><table><thead><tr>" + "".join(f"<th>{h}</th>" for h in head) + "</tr></thead><tbody>"
        det += "".join("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>" for r in body[:200])
        det += "</tbody></table>"
    return f"""<!DOCTYPE html><html lang="es"><head><meta charset="utf-8">
<title>Reporte {tipo} - FashionStore</title>
<style>body{{font-family:Arial,sans-serif;margin:24px;color:#0f172a}}
h1{{color:#e11d48}}table{{border-collapse:collapse;width:100%;margin-top:12px}}
th{{background:#e11d48;color:#fff;padding:8px;text-align:left}}
td{{border:1px solid #e2e8f0;padding:6px 8px;font-size:13px}}
footer{{margin-top:24px;color:#64748b;font-size:12px}}</style></head>
<body><h1>FashionStore — Reporte {tipo}</h1>
<p>Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}</p>
<h2>Resumen</h2><table><tbody>{resumen}</tbody></table>{det}
<footer>Sistema de Información FashionStore · Uso interno gerencial</footer>
</body></html>"""


def exportar_pdf(tipo: str, datos: Dict[str, Any]) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    els = [Paragraph(f"FashionStore — Reporte {tipo}", styles["Heading1"]),
           Paragraph(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles["Normal"]),
           Spacer(1, 12), Paragraph("Resumen", styles["Heading2"])]
    rdata = [["Campo", "Valor"]] + [[str(a), str(b)] for a, b in _resumen_rows(datos)]
    t = Table(rdata, colWidths=[220, 300])
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E11D48")),
                           ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                           ("FONTSIZE", (0, 0), (-1, -1), 9)]))
    els += [t, Spacer(1, 12)]
    head, body = _detalle_rows(tipo, datos)
    if head:
        els.append(Paragraph("Detalle (primeras 50 filas)", styles["Heading2"]))
        ddata = [head] + [[str(c) for c in r] for r in body[:50]]
        dt = Table(ddata, repeatRows=1)
        dt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
                                ("FONTSIZE", (0, 0), (-1, -1), 8)]))
        els.append(dt)
        if len(body) > 50:
            els += [Spacer(1, 6), Paragraph(f"… y {len(body) - 50} filas más (ver Excel/HTML).", styles["Normal"])]
    doc.build(els)
    return buf.getvalue()


FORMATO_MEDIA = {
    "CSV": ("text/csv; charset=utf-8", "csv"),
    "EXCEL": ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"),
    "HTML": ("text/html; charset=utf-8", "html"),
    "PDF": ("application/pdf", "pdf"),
}


def responder_reporte(tipo: str, datos: Dict[str, Any], formato: str):
    """Devuelve (content_bytes_or_str, media_type, filename). CSV lo resuelve el llamador legacy."""
    f = (formato or "JSON").upper()
    tag = _fecha_tag()
    if f == "EXCEL":
        return exportar_excel(tipo, datos), *FORMATO_MEDIA["EXCEL"][:1], f"reporte_{tipo.lower()}_{tag}.xlsx"
    if f == "HTML":
        return exportar_html(tipo, datos), *FORMATO_MEDIA["HTML"][:1], f"reporte_{tipo.lower()}_{tag}.html"
    if f == "PDF":
        return exportar_pdf(tipo, datos), *FORMATO_MEDIA["PDF"][:1], f"reporte_{tipo.lower()}_{tag}.pdf"
    return datos, "application/json", ""
